"""
End-to-end test of the new pipeline (Flow 1 + Flow 2) with no network:
FalkorDB Lite as the graph, FakeLLM for summaries, HashEmbedder for vectors.

    uv run pytest tests/test_codegraph_e2e.py -s
"""

import asyncio
import shutil
import subprocess
from pathlib import Path

import pytest

from github_graphrag.codegraph.config import FakeLLM, HashEmbedder, Settings
from github_graphrag.codegraph.pipeline import ingest
from github_graphrag.codegraph.retrieve import ask, impact
from github_graphrag.codegraph.sync import delete_repo

from falkor_store import FalkorStore

FIXTURE = Path(__file__).parent / "fixtures" / "mini_repo"
REPO = "mini"
SEC = "backend/app/core/security.py"


def git(root, *args):
    subprocess.run(["git", "-C", str(root), *args], check=True, capture_output=True)


def commit(root, msg):
    git(root, "add", "-A")
    git(root, "-c", "user.email=test@example.com", "-c", "user.name=test", "commit", "-qm", msg)


def make_deck(path: Path):
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Claim auto-approval"
    s.placeholders[1].text = ("Claims with all documents attached and a low fraud score are "
                              "approved automatically; risky claims go to manual review.")
    s.notes_slide.notes_text_frame.text = "Fraud score comes from the claims decision pipeline."
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


@pytest.fixture
def repo_root(tmp_path):
    root = tmp_path / "mini"
    shutil.copytree(FIXTURE, root)
    make_deck(root / "product" / "Claims_Overview.pptx")
    git(root, "init", "-q")
    commit(root, "initial")
    return root


def edges(store, rel):
    return {(r["a"].split(":", 1)[1], r["b"].split(":", 1)[1]) for r in store.run(
        f"MATCH (a)-[r:{rel}]->(b) RETURN a.id AS a, b.id AS b")}


def run(store, llm, emb, root, settings):
    return asyncio.run(ingest(store, llm, emb, root, REPO, settings))


def test_new_repo_then_updates(repo_root):
    store = FalkorStore()
    llm, emb = FakeLLM(), HashEmbedder()
    settings = Settings(llm_provider="fake", embedding_provider="hash",
                        bridge_min_score=0.55, bridge_top_k=2)

    # ================================================= Flow 1: new repo
    r1 = run(store, llm, emb, repo_root, settings)
    print("\nFLOW 1", r1)
    assert r1["mode"] == "new repo"
    assert r1["sync"]["added"] == 16            # 13 .py + 2 .md + 1 .pptx
    assert r1["enrich"]["llm_calls"] > 0 and r1["enrich"]["failed"] == 0

    calls = edges(store, "CALLS")
    assert ("backend/app/api/auth.py:login", f"{SEC}:verify_password") in calls
    assert ("backend/app/api/auth.py:login",
            "backend/app/db/users.py:UserRepository.find_by_email") in calls   # typed var
    assert ("backend/app/db/users.py:UserRepository.create",
            "backend/app/db/base.py:BaseRepository.save") in calls            # inheritance
    assert ("backend/app/db/users.py:UserRepository",
            "backend/app/db/base.py:BaseRepository") in edges(store, "INHERITS")

    mentions = edges(store, "MENTIONS")
    assert ("docs/auth.md#Authentication > Login", f"{SEC}:verify_password") in mentions
    assert ("docs/auth.md#Authentication > Login", "POST /auth/login") in mentions
    # the '# not a heading' line inside the code fence must not create a section
    assert not store.run("MATCH (s:Section) WHERE s.heading CONTAINS 'not a heading' RETURN s")

    deck = store.run("MATCH (s:Section)-[r:RELATES_TO]->(m:Module) "
                     "WHERE s.doc_type = 'product' RETURN m.name AS m, r.score AS score")
    print("deck -> modules", deck)
    assert any(d["m"] == "backend/app/claims" for d in deck)

    # ================================================= no changes -> no work
    r2 = run(store, llm, emb, repo_root, settings)
    print("NO-OP ", r2)
    assert r2["mode"] == "update"
    assert r2["sync"]["unchanged"] == 16
    assert r2["enrich"]["llm_calls"] == 0 and r2["embed"]["functions"] == 0
    assert r2["modules"]["summarised"] == 0

    # ================================================= Flow 2: code + doc changes
    sec = repo_root / SEC
    text = sec.read_text()
    text = text.replace(
        "return hmac.compare_digest(hash_password(plain), hashed)",
        "if not plain:\n        return False\n    return hmac.compare_digest(hash_password(plain), hashed)")
    text = text.replace(
        '\n\ndef _sign(payload: str) -> str:\n    return hmac.new(SECRET.encode(), payload.encode(), "sha256").hexdigest()\n',
        '\n\ndef rotate_secret(new_secret: str) -> None:\n    """Replace the signing secret."""\n'
        '    global SECRET\n    SECRET = new_secret\n')
    sec.write_text(text)
    (repo_root / "backend/app/api/claims.py").unlink()
    claims_doc = repo_root / "docs/claims.md"
    claims_doc.write_text(claims_doc.read_text() + "\nRejected claims can be appealed within 30 days.\n")
    commit(repo_root, "change security, drop claims api, edit docs")

    r3 = run(store, llm, emb, repo_root, settings)
    print("FLOW 2", r3)
    assert r3["sync"]["modified"] == 2 and r3["sync"]["deleted"] == 1
    # verify_password changed, rotate_secret new (create_token/_sign: _sign removed)
    assert 1 <= r3["enrich"]["llm_calls"] <= 3
    assert r3["sync"]["doc_links_flagged_outdated"] >= 1
    assert not store.run("MATCH (f:Function) WHERE f.name = '_sign' RETURN f")
    assert not store.run("MATCH (e:Endpoint {route: '/claims'}) RETURN e")
    assert store.run("MATCH (f:Function {name: 'rotate_secret'}) RETURN f")
    outdated = store.run("MATCH (s:Section)-[r]->(f:Function {name: 'verify_password'}) "
                         "RETURN r.doc_may_be_outdated AS flag")
    assert any(o["flag"] for o in outdated)

    # ================================================= impact
    res = impact(store, REPO, "hash_password")
    callers = {c["qname"] for c in res["callers"]}
    print("IMPACT hash_password", callers, res["endpoints"])
    assert {"verify_password", "register", "login"} <= callers
    assert {e["endpoint"] for e in res["endpoints"]} == {"POST /auth/login", "POST /auth/register"}

    # ================================================= ask
    a = asyncio.run(ask(store, llm, emb, REPO, "What breaks if I change `hash_password`?"))
    assert a.route == "impact" and "POST /auth/register" in a.context
    a = asyncio.run(ask(store, llm, emb, REPO, "How is the login password checked?"))
    assert a.route == "explain" and "verify_password" in a.context
    print("ASK sources", a.sources)

    # ================================================= cleanup
    assert delete_repo(store, REPO) > 0
    assert not store.run("MATCH (n) WHERE n.repo = $r RETURN n LIMIT 1", r=REPO)
